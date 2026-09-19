"""Build the webinar slide deck from the Xebia template.

This script is the only way the deck changes. The PPTX in `docs/` is generated from
the SLIDES list below and should never be edited by hand: change the data here,
rebuild, regenerate the PDF, and then update the matching heading in
PRESENTER_SCRIPT.md.

Run from the repository root:

    uv run --with python-pptx python instructor/build_deck.py
    uv run --with python-pptx python instructor/build_deck.py --inventory
    uv run --with python-pptx python instructor/build_deck.py --check-script

Then convert to PDF with LibreOffice, as described in the root CLAUDE.md.
"""

from __future__ import annotations

import argparse
import copy
import io
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "instructor" / "Xebia PPT General Template.pptx"
OUTPUT = ROOT / "docs" / "modernize-legacy-code-webinar-slides.pptx"
SCRIPT = ROOT / "instructor" / "PRESENTER_SCRIPT.md"

# Diagrams are drawn with the diagram-builder MCP server. Each .svg in this folder is
# the editable source, and the .png next to it is the export the deck embeds.
DIAGRAMS = ROOT / "instructor" / "diagrams"

# Layout indexes in the Xebia template, named so the slide list below reads clearly.
COVER = 0          # Main-cover_light
BASIC = 6          # Content_Basic
THREE_COL = 8      # Content_3 Columns_Comparison_2
TWO_COL = 9        # Content_2 Columns
BREAK = 24         # 1_Photo_background
END = 27           # End-cover_light

# The template's own body default is 12 pt, which is unreadable on a shared screen,
# so every run gets an explicit size instead of inheriting one.
BULLET_PT = 20
CODE_PT = 18
COLUMN_CODE_PT = 16
COLUMN_BODY_PT = 16
COLUMN_HEADING_PT = 16
SUBTITLE_PT = 16
MONO = "Consolas"

# Density rules from CLAUDE.md, checked against this data before anything is built.
MAX_SLIDES = 20
MAX_BULLETS = 6
MAX_WORDS = 11     # "under a dozen words"
MAX_CODE_LINES = 12

# Column placeholders aren't in left-to-right index order, so each column is named by
# its (heading idx, body idx) pair, measured from the template.
THREE_COL_SLOTS = [(11, 1), (17, 16), (14, 13)]    # left, middle, right
TWO_COL_SLOTS = [(11, 1), (14, 13)]                # left, right

# The two-column layout reserves a band for icon pictures above each column. We don't
# use icons, and once those frames are removed the columns sit alone in the bottom half
# of the slide. Moving them up to the three-column layout's positions closes the gap.
TWO_COL_HEADING_TOP_IN = 2.16
TWO_COL_BODY_TOP_IN = 2.84
TWO_COL_BODY_HEIGHT_IN = 4.06

# Diagram slides fill the body area of the basic layout, under the title.
DIAGRAM_LEFT_IN = 0.61
DIAGRAM_TOP_IN = 1.85
DIAGRAM_MAX_HEIGHT_IN = 5.0
DIAGRAM_MAX_WIDTH_IN = 12.1


# ---------------------------------------------------------------------------- slides
#
# Text wrapped in backticks is set in a monospace font, so method names, file paths
# and code read as code. Everything else is set in the template's Arial.

SLIDES = [
    {
        "kind": "Title",
        "layout": COVER,
        # The line break keeps the title clear of the cover's artwork.
        "title": "Modernizing Legacy\nCode with AI",
        "subtitle": ("A Practical Approach to Brownfield Projects\n"
                     "Eric Greene · Xebia\neric.greene@xebia.com"),
        "notes": (
            "Welcome and introductions. Ninety minutes: about twenty of slides and seventy of "
            "live demos in VS Code with GitHub Copilot. Questions are welcome at any time, in "
            "the chat or out loud. Everything is in the webinar repo. Set the frame: most of "
            "us spend our days in code someone else wrote, and that's where AI tools are "
            "hardest to use well."
        ),
    },
    {
        "kind": "Agenda",
        "layout": BASIC,
        "title": "Agenda",
        "bullets": [
            "Why brownfield is harder for AI",
            "Using AI to understand unfamiliar code",
            "Building a safety net before refactoring",
            "Strangler fig, targeted refactors, and dependency upgrades",
            "Translating, untangling, and rewriting while preserving behavior",
            "Knowing when to trust, intervene, or leave alone",
        ],
        "notes": (
            "Six topics, matching the README agenda, and four demos on one legacy codebase, "
            "ShipBill, a courier's billing system from about 2012. Topics one and two lead into "
            "Demo 1, topic three into Demo 2, topic four into Demos 3 and 4, and topics five "
            "and six into Demo 4. Each demo starts where the last one stopped."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Why brownfield is harder for AI",
        "columns": [
            ("Missing context", [
                "The reasons left with the people",
                "Docs and comments go stale",
            ]),
            ("Implicit constraints", [
                "Other systems need exact output",
                "Quirks already on every invoice",
            ]),
            ("Accumulated tech debt", [
                "God functions and global state",
                "Old dependencies, few tests",
            ]),
        ],
        "notes": (
            "On a greenfield project the agent sees everything that matters. In brownfield, "
            "the most important facts aren't in the code: why customer 1047 gets 15 percent, "
            "which downstream import reads the output, which rounding quirk is now part of "
            "every invoice. Tech debt makes it worse, because big functions and globals give "
            "the agent no safe place to change one thing. The agent will still sound "
            "confident, so our job is to supply the context and the checks."
        ),
    },
    {
        "kind": "Diagram",
        "layout": BASIC,
        "title": "A loop for brownfield work",
        "image": "brownfield-loop.png",
        # Trims the diagram page's empty top and bottom margins, as fractions.
        "crop": (0.27, 0.2),
        "notes": (
            "This is the shape of the whole session. Understand the code, protect its behavior, "
            "change one small thing, verify, and go round again. In every step the AI does "
            "the heavy lifting, and a person does the part that needs judgment: checking "
            "claims, approving what counts as unchanged, reviewing the diff, and deciding "
            "what to do about a difference. The demos follow the loop in order."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "Using AI to understand unfamiliar code",
        "bullets": [
            "Run it first, and read the real output",
            "Scope the agent to code, not old docs",
            "Map the architecture, then trace one real calculation",
            "Ask for `file:line` evidence for every claim",
            "Ask what the code can't tell you",
            "Generate docs, then review them like code",
        ],
        "notes": (
            "Start with the running system, because it's the only source that can't be out "
            "of date. Tell the agent to work from the code, since old READMEs will mislead it. "
            "Mapping gives you the shape; tracing one real number end to end gives you the "
            "details that matter. Citations make claims checkable. Asking what the code "
            "can't explain produces the list of questions for the business. Generated docs "
            "are a draft until someone checks them."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "Verify the map, then keep it",
        "columns": [
            ("Check the claims", [
                "A mechanical map: imports, callers, dead code",
                "Open the lines the agent cites",
                "Run the code to confirm a rule",
            ]),
            ("Capture the context", [
                "Reviewed architecture notes for people",
                "Custom instructions for every AI session",
                "Open questions with named owners",
            ]),
        ],
        "notes": (
            "Anything you can check mechanically, check mechanically. Demo 1 uses a small "
            "ast-based script to confirm the import graph and find dead code, then opens cited "
            "lines and runs a billing month to confirm a rule. Then write it down. Architecture "
            "notes are for people, custom instructions are for every future Copilot session, "
            "and open questions go to named owners in the business."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 1: Read before you change",
        "subtitle": "demos/01-understand-legacy · about 17 minutes",
        "notes": (
            "Run billing for September with customer 1047, read the outdated README.txt, glance "
            "at calc_invoice. Ask role: architecture, trace 1047's total, unexplained rules. Run "
            "codemap.py and compare, including the dead fax code. Show September missing from "
            "FUEL_TABLE and the fuel_rate fallback. Agent role: /document-legacy, review three citations, then "
            "/capture-context."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "Characterization tests",
        "bullets": [
            "Pin what the code does, not what it should",
            "Expected values come from running the real code",
            "Pin the quirks too, and label them",
            "Find a seam to test without changing code",
            "Bugs become open questions, not fixes",
        ],
        "notes": (
            "Michael Feathers' term. A characterization test records today's behavior so you'll "
            "notice when it changes. The expected value comes from running the code, never from "
            "the agent's arithmetic or anyone's idea of correct. Quirks get pinned with a "
            "comment saying it's deliberate. A seam is a place to change behavior from outside, "
            "like ShipBill's database setting. Bugs found now go on the open-questions list, "
            "because fixing them mid-refactor hides what changed."
        ),
    },
    {
        "kind": "Code",
        "layout": BASIC,
        "title": "Pinning a quirk, on purpose",
        "code": [
            "# tests/test_util.py",
            "def test_round2_rounds_down_a_half_cent_produced_by_float_math():",
            "    # 6.95 + 0.95 * 8.5 is 15.024999... in floats,",
            "    # so customers were billed 15.02, not 15.03.",
            '    assert util.round2(rates.base_charge("ECO", 4, 9.5)) == 15.02',
        ],
        "notes": (
            "A real test from the demos. Exact arithmetic says this parcel costs 15.025, which "
            "rounds to 15.03. The legacy float math lands a hair below and bills 15.02, and "
            "that's what customers have paid for years. The test pins it and the comment says "
            "it's deliberate. Note the exact equality: no approx, because a one-cent change is "
            "the whole point. This exact quirk comes back in Demo 4."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "Two layers of safety net",
        "columns": [
            ("Golden master", [
                "A whole month's output, compared byte for byte",
                "Fast to build, but only says something changed",
                "A person reviews it before it's approved",
            ]),
            ("Focused characterization tests", [
                "One rule per test, named after the rule",
                "A failure names the rule that broke",
                "Coverage shows the gaps, not the safety",
            ]),
        ],
        "notes": (
            "A golden master runs the system end to end and compares the output with a reviewed, "
            "approved copy. In Demo 2 four tests reach 88 percent coverage, but a failure only "
            "tells you something moved. Focused tests tell you which rule. Coverage is a map of "
            "what ran: the one skipped legacy test still reports 18 percent. Use it to find gaps "
            "worth closing, not as a target."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 2: Build the safety net",
        "subtitle": "demos/02-safety-net · about 18 minutes",
        "notes": (
            "Coverage of the skipped test: 18 percent. Read conftest.py, the seam, and "
            "test_golden_master.py. Run pytest, review tests/received, run approve.py, rerun. "
            "Coverage: 88 percent. Read the prompt and instruction files, then /characterize "
            "rates.py and billing.py. Close the weekend due-date gap. Break RETURN_RATE and "
            "watch both layers fail, then put it back."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Incremental modernization patterns",
        "columns": [
            ("Targeted refactor", [
                "Change structure, never behavior",
                "One step, tests before and after",
            ]),
            ("Dependency upgrade", [
                "Warnings show what will break",
                "One package, fix by cause",
            ]),
            ("Strangler fig", [
                "Grow new code behind a seam",
                "Move one slice at a time",
            ]),
        ],
        "notes": (
            "Three patterns, all small steps under green tests. A targeted refactor reshapes "
            "code without changing what it does, so the tests prove which one you did. A "
            "dependency upgrade starts from the warnings you already have, one dependency at a "
            "time. The strangler fig replaces a component slice by slice behind a seam, with "
            "no big switchover. Demo 3 does the first two; Demo 4 does the third."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "Refactors and upgrades with an agent",
        "bullets": [
            "Plan first, and read the plan for behavior changes",
            "Prompt files make the safe workflow repeatable",
            "The agent runs the tests; you review each diff",
            "Never let the agent edit a test to pass",
            "Group failures by cause before fixing anything",
            "Finish with warnings treated as errors",
        ],
        "notes": (
            "The Plan role is where you catch an agent that wants to tidy up the rules while it "
            "refactors. Prompt files package the safe workflow: run the tests, make one change, "
            "undo if it fails. The agent does the running; you do the reviewing. The one rule "
            "that's never negotiable is that tests don't change to make code pass. For upgrades, "
            "group failures by cause; Demo 3's 32 failures come from one removed feature."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 3: Modernize in small steps",
        "subtitle": "demos/03-small-steps · about 17 minutes",
        "notes": (
            "72 tests green, point at the arrow warnings. Plan role for extracting "
            "price_shipment, then /refactor-step in the Agent role; rerun tests and billing, total "
            "still 3,348.55. Then pytest with deprecations as errors: 32 fail. uv add "
            "arrow==1.4.0, same 32 fail differently. /upgrade-dependency, watch for the "
            "timestamp second wave. Finish green with warnings as errors."
        ),
    },
    {
        "kind": "Diagram",
        "layout": BASIC,
        "title": "The strangler fig",
        "image": "strangler-fig.png",
        "crop": (0.15, 0.13),
        "notes": (
            "Named for a fig that grows down around its host tree and eventually replaces it. "
            "calc_invoice asks the router for a price and never knows which engine answered. "
            "Services in ROUTED_TO_V2 go to the new engine, everything else to the legacy one. "
            "The parallel run prices every shipment both ways before a service moves. One "
            "service at a time, and when nothing routes to the old code, delete it."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Translate, untangle, or rewrite",
        "columns": [
            ("Translate", [
                "New language, same results",
                "Recorded runs become the test",
            ]),
            ("Untangle", [
                "Find seams, extract pure functions",
                "Same arithmetic, same order",
            ]),
            ("Rewrite", [
                "New code beside the old code",
                "A parallel run on real data decides",
            ]),
        ],
        "notes": (
            "Three workflows, one principle: behavior is preserved and proved. Translating "
            "COBOL, VB, or old Java to something modern works like a golden master across "
            "languages: record what the old system does and replay it. Untangling is Demo 3's "
            "refactor. Rewriting is Demo 4: the new engine lives beside the old one, and a "
            "parallel run on every real shipment decides when it's ready."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Knowing the limits",
        "columns": [
            ("Trust the agent", [
                "Reading, mapping, drafting tests",
                "Changes under a strong safety net",
            ]),
            ("Intervene", [
                "Even a one-cent change",
                "Edits to tests or external contracts",
            ]),
            ("Leave it alone", [
                "Works, stable, has a contract",
                "Guard it with a golden master",
            ]),
        ],
        "notes": (
            "Let the agent run where mistakes are cheap and caught fast. Step in when the "
            "evidence shows a behavior change, when the agent wants to edit a test, or when a "
            "change touches another system's contract; those are people's decisions, and they "
            "get written down. Some code should be left alone: ShipBill's export works, rarely "
            "changes, and a mistake fails accounting's import. Demo 4 shows all three."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 4: Strangle, rewrite, and know when to stop",
        "subtitle": "demos/04-strangler-fig · about 18 minutes",
        "notes": (
            "Read pricing_router.py, the changed call in billing.py, and parallel_run.py. "
            "/strangle-pricing, review pricing_v2.py, read the parallel run: about 15 one-cent "
            "disagreements. Ask why: float math. Decide to keep legacy cents, record D-002, "
            "rerun to 272 of 272. Route ECO, tests and billing unchanged. Ask to modernize "
            "export.py and watch Copilot push back."
        ),
    },
    {
        "kind": "Wrap",
        "layout": BASIC,
        "title": "Next steps",
        "bullets": [
            "Clone the repo and run the four demos",
            "Pick one legacy module and map it first",
            "Put a golden master around it before changing it",
            "Write your team's rules into custom instructions",
            "Adapt the repo's modernization checklist for your team",
        ],
        "notes": (
            "Point people to demos/README.md for setup. Suggest starting small: one module, "
            "mapped and documented, with a golden master around it, before any change. The "
            "checklist in docs/legacy-modernization-checklist.md follows the same loop as the "
            "session. Mention the Xebia Academy link in the README too."
        ),
    },
    {
        "kind": "Closing",
        "layout": END,
        "title": "Thank you",
        "subtitle": "Eric Greene · Xebia\neric.greene@xebia.com",
        "notes": (
            "Thank the audience. Leave this slide up while you answer any final questions, and "
            "repeat that the slides, demos, and checklist are in the repo."
        ),
    },
]


# ---------------------------------------------------------------------- validation

def words(text: str) -> int:
    """Count words the way a reader would: tokens containing a letter or digit."""
    return sum(1 for t in text.replace("`", "").split() if re.search(r"[A-Za-z0-9]", t))


def validate_data(slides: list[dict]) -> list[str]:
    problems = []
    if len(slides) > MAX_SLIDES:
        problems.append(f"{len(slides)} slides, limit is {MAX_SLIDES}")
    for n, s in enumerate(slides, start=1):
        groups = []
        if "bullets" in s:
            groups.append(s["bullets"])
        for _, items in s.get("columns", []):
            groups.append(items)
        total = sum(len(g) for g in groups)
        if total > MAX_BULLETS:
            problems.append(f"slide {n}: {total} bullets, limit is {MAX_BULLETS}")
        for g in groups:
            for b in g:
                if words(b) > MAX_WORDS:
                    problems.append(f"slide {n}: {words(b)} words in {b!r}")
        code_blocks = [s.get("code", [])] + [lines for _, lines in s.get("code_columns", [])]
        for block in code_blocks:
            if len(block) > MAX_CODE_LINES:
                problems.append(f"slide {n}: {len(block)} code lines, limit is {MAX_CODE_LINES}")
        if "image" in s and not (DIAGRAMS / s["image"]).is_file():
            problems.append(f"slide {n}: diagram {s['image']} is missing; export it first")
        if not s.get("notes", "").strip():
            problems.append(f"slide {n}: speaker notes are empty")
    for n in (1, len(slides)):
        if "eric.greene@xebia.com" not in slides[n - 1].get("subtitle", ""):
            problems.append(f"slide {n}: the presenter's email must be on the first and last slides")
    return problems


# ---------------------------------------------------------------------------- build

def add_runs(paragraph, text: str, size: int) -> None:
    """Split on backticks so `code` segments get the monospace font."""
    for i, part in enumerate(text.split("`")):
        if not part:
            continue
        run = paragraph.add_run()
        run.text = part
        run.font.size = Pt(size)
        if i % 2 == 1:
            run.font.name = MONO


def fill_lines(placeholder, lines: list[str], size: int) -> None:
    """Write one paragraph per line, without leaving a stray empty first paragraph."""
    tf = placeholder.text_frame
    tf.clear()
    first = True
    for line in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        add_runs(para, line, size)


def fill_code(placeholder, lines: list[str], size: int) -> None:
    """Write code as unbulleted monospace lines.

    The body placeholder bullets every paragraph, which would put a dot in front of
    each line of code, so each paragraph gets an explicit buNone and no indent.
    """
    tf = placeholder.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_pr = para._p.get_or_add_pPr()
        p_pr.set("marL", "0")
        p_pr.set("indent", "0")
        p_pr.insert(0, p_pr.makeelement(qn("a:buNone"), {}))
        run = para.add_run()
        # A non-breaking space keeps blank lines and leading indentation intact.
        run.text = line.replace(" ", " ") if line else " "
        run.font.size = Pt(size)
        run.font.name = MONO


def transparent_png(path: Path) -> io.BytesIO:
    """Make the diagram's white page background transparent.

    The diagram tool exports on solid white, which would cover the template's light
    background artwork. Pillow comes with python-pptx, so no extra dependency is needed.
    """
    img = Image.open(path).convert("RGBA")
    img.putdata([(r, g, b, 0) if min(r, g, b) >= 250 else (r, g, b, a)
                 for r, g, b, a in img.get_flattened_data()])
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def add_diagram(slide, image: str, crop: tuple[float, float]) -> None:
    """Place a diagram under the title, scaled to the body area and cropped.

    The diagram tool always exports a full landscape page, so the empty margins above
    and below the drawing are cropped off before the picture is sized.
    """
    top_crop, bottom_crop = crop
    pic = slide.shapes.add_picture(transparent_png(DIAGRAMS / image), Inches(DIAGRAM_LEFT_IN),
                                   Inches(DIAGRAM_TOP_IN))
    pic.crop_top, pic.crop_bottom = top_crop, bottom_crop
    aspect = pic.image.size[0] / (pic.image.size[1] * (1 - top_crop - bottom_crop))
    height = min(DIAGRAM_MAX_HEIGHT_IN, DIAGRAM_MAX_WIDTH_IN / aspect)
    pic.height, pic.width = Inches(height), Inches(height * aspect)


def move(placeholder, *, top: float, height: float | None = None) -> None:
    """Reposition a placeholder vertically, keeping its inherited left and width.

    A slide placeholder inherits its geometry from the layout until it's given its own,
    and python-pptx needs all four values set to write one, so the current left and
    width are carried over explicitly.
    """
    left, width = placeholder.left, placeholder.width
    new_height = Inches(height) if height is not None else placeholder.height
    placeholder.left, placeholder.width = left, width
    placeholder.top, placeholder.height = Inches(top), new_height


def restore_slide_number(slide) -> None:
    """python-pptx skips slide-number placeholders when cloning a layout, but the
    template's own slides carry them, so copy the layout's one across."""
    for lp in slide.slide_layout.placeholders:
        if lp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(lp._element))


def remove_unfilled(slide, filled: set[int]) -> None:
    """An unfilled picture placeholder renders as an empty frame, so anything we
    didn't write to is removed rather than left behind."""
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in filled:
            ph._element.getparent().remove(ph._element)


def build(output: Path) -> None:
    prs = Presentation(TEMPLATE)
    layouts = prs.slide_masters[0].slide_layouts
    assert len(layouts) == 29, f"template has {len(layouts)} layouts, expected 29"

    # Remove the template's example slides. Layouts, the master and the notes master
    # survive, which is all the deck needs.
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    id_list = prs.slides._sldIdLst
    for sld in list(id_list):
        prs.part.drop_rel(sld.get(rel_id))
        id_list.remove(sld)
    assert len(prs.slides) == 0

    for s in SLIDES:
        slide = prs.slides.add_slide(layouts[s["layout"]])
        ph = {p.placeholder_format.idx: p for p in slide.placeholders}
        filled = {0}
        ph[0].text_frame.text = s["title"]

        if s["layout"] in (COVER, BREAK, END):
            fill_lines(ph[1], s["subtitle"].split("\n"), SUBTITLE_PT)
            filled.add(1)
        elif "columns" in s:
            slots = THREE_COL_SLOTS if s["layout"] == THREE_COL else TWO_COL_SLOTS
            for (head_idx, body_idx), (heading, items) in zip(slots, s["columns"]):
                fill_lines(ph[head_idx], [heading], COLUMN_HEADING_PT)
                fill_lines(ph[body_idx], items, COLUMN_BODY_PT)
                filled.update({head_idx, body_idx})
                if s["layout"] == TWO_COL:
                    move(ph[head_idx], top=TWO_COL_HEADING_TOP_IN)
                    move(ph[body_idx], top=TWO_COL_BODY_TOP_IN,
                         height=TWO_COL_BODY_HEIGHT_IN)
        elif "code_columns" in s:
            for (head_idx, body_idx), (heading, lines) in zip(TWO_COL_SLOTS, s["code_columns"]):
                fill_lines(ph[head_idx], [heading], COLUMN_HEADING_PT)
                fill_code(ph[body_idx], lines, COLUMN_CODE_PT)
                filled.update({head_idx, body_idx})
                move(ph[head_idx], top=TWO_COL_HEADING_TOP_IN)
                move(ph[body_idx], top=TWO_COL_BODY_TOP_IN, height=TWO_COL_BODY_HEIGHT_IN)
        elif "code" in s:
            fill_code(ph[1], s["code"], CODE_PT)
            filled.add(1)
        elif "image" in s:
            add_diagram(slide, s["image"], s["crop"])
        else:
            fill_lines(ph[1], s["bullets"], BULLET_PT)
            filled.add(1)

        remove_unfilled(slide, filled)
        if s["layout"] not in (COVER, END):
            restore_slide_number(slide)
        slide.notes_slide.notes_text_frame.text = s["notes"]

    assert len(prs.slides) == len(SLIDES)
    for n, slide in enumerate(prs.slides, start=1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f"slide {n} has no notes"

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def inventory() -> str:
    rows = [f"{'#':>2}  {'kind':<8} title"]
    for n, s in enumerate(SLIDES, start=1):
        title = s["title"].replace("\n", " ")
        rows.append(f"{n:>2}  {s['kind']:<8} {title}")
    return "\n".join(rows)


def check_script() -> list[str]:
    """Compare the presenter script's slide headings with the deck, line for line.

    Each slide in PRESENTER_SCRIPT.md has a heading of the form "## Slide N: Title",
    and the titles drift whenever a slide is renamed without updating the script.
    """
    expected = [f"## Slide {n}: {s['title'].replace(chr(10), ' ')}"
                for n, s in enumerate(SLIDES, start=1)]
    text = SCRIPT.read_text(encoding="utf-8") if SCRIPT.exists() else ""
    found = [line.rstrip() for line in text.splitlines() if line.startswith("## Slide ")]
    problems = []
    for i in range(max(len(expected), len(found))):
        want = expected[i] if i < len(expected) else "(no slide)"
        got = found[i] if i < len(found) else "(missing)"
        if want != got:
            problems.append(f"expected {want!r}, found {got!r}")
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inventory", action="store_true",
                        help="print the slide list and exit, for checking the script against it")
    parser.add_argument("--check-script", action="store_true",
                        help="check PRESENTER_SCRIPT.md's slide headings against the deck and exit")
    args = parser.parse_args()

    if args.inventory:
        print(inventory())
        return 0

    if args.check_script:
        problems = check_script()
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
        print("presenter script matches the deck" if not problems
              else f"{len(problems)} heading(s) don't match", file=sys.stderr if problems else sys.stdout)
        return 1 if problems else 0

    problems = validate_data(SLIDES)
    if problems:
        print("Slide data breaks the rules in CLAUDE.md:", file=sys.stderr)
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
        return 1

    build(OUTPUT)
    print(f"built {OUTPUT.relative_to(ROOT)} with {len(SLIDES)} slides")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
